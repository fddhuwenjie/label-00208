"""文档解析服务 - 支持 docx, doc, xlsx, xls, pdf, txt

文档格式支持说明：
- .docx: 完整支持（python-docx）
- .doc: 多级降级支持（python-docx → antiword → textract）
        自动尝试多种解析方式，最大程度兼容旧版 Word 格式
- .xlsx: 完整支持（openpyxl via pandas）
- .xls: 完整支持（xlrd via pandas），旧版 Excel 97-2003 格式
- .pdf: 完整支持（pdfplumber）
- .txt: 完整支持（chardet 自动编码检测）
"""

import os
import subprocess
from pathlib import Path
from datetime import datetime
from typing import Optional
import chardet

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
import pandas as pd
import pdfplumber

from app.models.document import Document, DocumentParagraph, ExcelDocument, ExcelSheet
from app.utils.file_utils import get_file_id, get_file_size
from app.utils.text_utils import clean_text, normalize_whitespace
from app.utils.logger import parser_logger as logger


class DocumentParser:
    """文档解析器"""
    
    def __init__(self):
        self.parsers = {
            ".docx": self._parse_docx,
            ".doc": self._parse_doc,
            ".xlsx": self._parse_excel,
            ".xls": self._parse_excel,
            ".pdf": self._parse_pdf,
            ".txt": self._parse_txt,
        }
    
    def parse(self, file_path: str) -> Optional[Document]:
        """
        解析文档
        
        Args:
            file_path: 文件路径
            
        Returns:
            Document 对象，解析失败返回 None
        """
        path = Path(file_path)
        
        if not path.exists():
            return None
        
        ext = path.suffix.lower()
        parser = self.parsers.get(ext)
        
        if parser is None:
            return None
        
        try:
            return parser(file_path)
        except Exception as e:
            logger.error(f"解析文件失败 {file_path}: {e}")
            return None
    
    def _create_base_document(self, file_path: str, file_type: str) -> Document:
        """创建基础文档对象"""
        path = Path(file_path)
        return Document(
            id=get_file_id(file_path),
            file_path=str(path.absolute()),
            file_name=path.name,
            file_type=file_type,
            file_size=get_file_size(file_path),
            created_at=datetime.now(),
        )
    
    def _parse_doc(self, file_path: str) -> Optional[Document]:
        """
        解析旧版 Word 文档（.doc 格式）
        
        处理策略：
        1. 首先尝试用 python-docx 解析（适用于 Office 2007+ 另存为的 .doc）
        2. 如果失败，尝试使用 antiword 命令行工具（需要系统安装）
        3. 如果 antiword 不可用，尝试使用 textract（需要额外安装）
        4. 最终降级：记录警告，返回 None
        
        注意：真正的 Word 97-2003 二进制格式需要 antiword 或 libreoffice 转换
        """
        # 策略1: 尝试用 docx 解析（某些 .doc 实际是 docx 格式）
        result = self._parse_docx(file_path)
        if result:
            return result
        
        # 策略2: 尝试使用 antiword（如果系统已安装）
        try:
            process = subprocess.run(
                ["antiword", file_path],
                capture_output=True,
                text=True,
                timeout=30
            )
            if process.returncode == 0 and process.stdout.strip():
                return self._create_doc_from_text(file_path, process.stdout)
        except FileNotFoundError:
            logger.debug("antiword 未安装，跳过 antiword 解析")
        except subprocess.TimeoutExpired:
            logger.warning(f"antiword 解析超时: {file_path}")
        except Exception as e:
            logger.debug(f"antiword 解析失败: {e}")
        
        # 策略3: 尝试使用 textract（如果已安装）
        try:
            import textract
            text = textract.process(file_path).decode("utf-8")
            if text.strip():
                return self._create_doc_from_text(file_path, text)
        except ImportError:
            logger.debug("textract 未安装，跳过 textract 解析")
        except Exception as e:
            logger.debug(f"textract 解析失败: {e}")
        
        logger.warning(
            f"无法解析 .doc 文件: {file_path}。"
            f"该文件可能是旧版 Word 97-2003 格式，建议：\n"
            f"1. 用 Word/WPS 将其另存为 .docx 格式\n"
            f"2. 或安装 antiword (brew install antiword / apt install antiword)"
        )
        return None
    
    def _create_doc_from_text(self, file_path: str, text: str) -> Document:
        """从纯文本创建文档对象"""
        doc = self._create_base_document(file_path, "doc")
        
        paragraphs = []
        full_text_parts = []
        
        para_texts = text.split("\n\n")
        for idx, para_text in enumerate(para_texts):
            para_text = clean_text(para_text)
            if not para_text:
                continue
            
            dp = DocumentParagraph(
                content=para_text,
                paragraph_index=idx,
            )
            paragraphs.append(dp)
            full_text_parts.append(para_text)
        
        doc.paragraphs = paragraphs
        doc.full_text = "\n\n".join(full_text_parts)
        return doc

    def _parse_docx(self, file_path: str) -> Optional[Document]:
        """解析 Word 文档（.docx 格式）"""
        try:
            docx = DocxDocument(file_path)
        except PackageNotFoundError:
            return None
        except Exception:
            return None
        
        doc = self._create_base_document(file_path, "docx")
        
        paragraphs = []
        full_text_parts = []
        
        for idx, para in enumerate(docx.paragraphs):
            text = clean_text(para.text)
            if not text:
                continue
            
            # 检测标题级别
            heading_level = None
            if para.style and para.style.name:
                style_name = para.style.name.lower()
                if "heading" in style_name or "标题" in style_name:
                    # 尝试提取标题级别
                    for i in range(1, 7):
                        if str(i) in style_name:
                            heading_level = i
                            break
                    if heading_level is None:
                        heading_level = 1
            
            dp = DocumentParagraph(
                content=text,
                paragraph_index=idx,
                heading_level=heading_level,
            )
            paragraphs.append(dp)
            full_text_parts.append(text)
        
        doc.paragraphs = paragraphs
        doc.full_text = "\n\n".join(full_text_parts)
        
        # 提取元数据
        try:
            core_props = docx.core_properties
            doc.title = core_props.title or None
            doc.author = core_props.author or None
        except Exception:
            pass
        
        return doc
    
    def _parse_excel(self, file_path: str) -> Optional[ExcelDocument]:
        """
        解析 Excel 文档（.xlsx 和 .xls 格式）
        
        处理策略：
        - .xlsx: 使用 openpyxl 引擎（完整支持）
        - .xls: 使用 xlrd 引擎（旧版格式支持）
        """
        path = Path(file_path)
        ext = path.suffix.lower()
        
        # 根据文件格式选择引擎
        engine = None
        if ext == ".xlsx":
            engine = "openpyxl"
        elif ext == ".xls":
            engine = "xlrd"
        
        try:
            xlsx = pd.ExcelFile(file_path, engine=engine)
        except ImportError as e:
            if "xlrd" in str(e):
                logger.warning(
                    f"无法解析 .xls 文件: {file_path}。"
                    f"xlrd 未安装，请运行 pip install xlrd 或将文件另存为 .xlsx 格式"
                )
            return None
        except Exception as e:
            logger.warning(f"解析 Excel 文件失败 {file_path}: {e}")
            return None
        
        file_type = "xlsx" if ext == ".xlsx" else "xls"
        doc = ExcelDocument(
            id=get_file_id(file_path),
            file_path=str(path.absolute()),
            file_name=path.name,
            file_type=file_type,
            file_size=get_file_size(file_path),
            created_at=datetime.now(),
        )
        
        sheets = []
        all_text_parts = []
        
        for sheet_name in xlsx.sheet_names:
            try:
                df = pd.read_excel(xlsx, sheet_name=sheet_name, header=None)
                
                # 转换为字符串列表
                rows = []
                for _, row in df.iterrows():
                    row_data = []
                    for cell in row:
                        if pd.notna(cell):
                            row_data.append(str(cell).strip())
                        else:
                            row_data.append("")
                    if any(row_data):  # 非空行
                        rows.append(row_data)
                
                sheet = ExcelSheet(name=sheet_name, rows=rows)
                sheets.append(sheet)
                
                # 构建文本
                all_text_parts.append(f"[工作表: {sheet_name}]")
                all_text_parts.append(sheet.to_text())
                
            except Exception as e:
                logger.warning(f"解析工作表 {sheet_name} 失败: {e}")
                continue
        
        doc.sheets = sheets
        doc.full_text = "\n\n".join(all_text_parts)
        
        # 创建段落（每行作为一个段落）
        paragraph_idx = 0
        for sheet in sheets:
            for row in sheet.rows:
                text = " | ".join(cell for cell in row if cell)
                if text.strip():
                    doc.paragraphs.append(DocumentParagraph(
                        content=text,
                        paragraph_index=paragraph_idx,
                    ))
                    paragraph_idx += 1
        
        return doc
    
    def _parse_pdf(self, file_path: str) -> Optional[Document]:
        """解析 PDF 文档"""
        doc = self._create_base_document(file_path, "pdf")
        
        paragraphs = []
        full_text_parts = []
        paragraph_idx = 0
        
        try:
            with pdfplumber.open(file_path) as pdf:
                doc.page_count = len(pdf.pages)
                
                for page_num, page in enumerate(pdf.pages, start=1):
                    try:
                        text = page.extract_text()
                        if not text:
                            continue
                        
                        # 按段落分割
                        page_paragraphs = text.split("\n\n")
                        
                        for para_text in page_paragraphs:
                            para_text = normalize_whitespace(para_text)
                            if not para_text or len(para_text) < 3:
                                continue
                            
                            dp = DocumentParagraph(
                                content=para_text,
                                paragraph_index=paragraph_idx,
                                page_number=page_num,
                            )
                            paragraphs.append(dp)
                            full_text_parts.append(para_text)
                            paragraph_idx += 1
                            
                    except Exception as e:
                        logger.warning(f"解析 PDF 第 {page_num} 页失败: {e}")
                        continue
                        
        except Exception as e:
            logger.error(f"打开 PDF 失败: {e}")
            return None
        
        doc.paragraphs = paragraphs
        doc.full_text = "\n\n".join(full_text_parts)
        
        return doc
    
    def _parse_txt(self, file_path: str) -> Optional[Document]:
        """解析纯文本文件"""
        doc = self._create_base_document(file_path, "txt")
        
        # 检测编码
        with open(file_path, "rb") as f:
            raw_data = f.read()
            detected = chardet.detect(raw_data)
            encoding = detected.get("encoding", "utf-8") or "utf-8"
        
        try:
            with open(file_path, "r", encoding=encoding, errors="replace") as f:
                content = f.read()
        except Exception:
            # 回退到 UTF-8
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        
        content = normalize_whitespace(content)
        
        # 按段落分割
        paragraphs = []
        para_texts = content.split("\n\n")
        
        for idx, para_text in enumerate(para_texts):
            para_text = para_text.strip()
            if not para_text:
                continue
            
            dp = DocumentParagraph(
                content=para_text,
                paragraph_index=idx,
            )
            paragraphs.append(dp)
        
        doc.paragraphs = paragraphs
        doc.full_text = content
        
        return doc


class DocumentScanner:
    """文档扫描器 - 扫描目录并解析所有文档"""
    
    def __init__(self):
        self.parser = DocumentParser()
    
    def scan_directory(
        self,
        directory: str,
        recursive: bool = True,
        progress_callback: callable = None
    ) -> list[Document]:
        """
        扫描目录解析所有文档
        
        Args:
            directory: 目录路径
            recursive: 是否递归扫描
            progress_callback: 进度回调函数 (current, total, file_name)
            
        Returns:
            解析成功的文档列表
        """
        from app.utils.file_utils import scan_directory as scan_dir
        
        # 收集所有文件
        files = list(scan_dir(directory, recursive=recursive))
        total = len(files)
        
        documents = []
        
        for idx, file_path in enumerate(files):
            if progress_callback:
                progress_callback(idx + 1, total, Path(file_path).name)
            
            doc = self.parser.parse(file_path)
            if doc:
                documents.append(doc)
        
        return documents
    
    def scan_single_file(self, file_path: str) -> Optional[Document]:
        """扫描单个文件"""
        return self.parser.parse(file_path)
